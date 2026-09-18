from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from app.rag.chunkers import ChunkingConfig, StructureAwareChunker
from app.rag.ingestion import IngestionService
from app.rag.parsers import DocumentParserRegistry, split_frontmatter
from app.rag.schemas import RagDocument


class FakeStore:
    def __init__(self) -> None:
        self.hashes = {}
        self.documents = {}
        self.replace_calls = 0

    def get_document_hash(self, doc_id):
        return self.hashes.get(doc_id)

    def replace_document(self, doc_id, chunks):
        self.replace_calls += 1
        self.documents[doc_id] = list(chunks)
        self.hashes[doc_id] = chunks[0].document_hash

    def delete_document(self, doc_id):
        chunks = self.documents.pop(doc_id, [])
        self.hashes.pop(doc_id, None)
        return len(chunks)


class RagPipelineTests(unittest.TestCase):
    def test_fenced_frontmatter_and_markdown_parser(self):
        raw = """```markdown
---
doc_id: card-a
category: product
tags: 联通, 四川
---
```

# **测试套餐**

## 资费

月租 30 元。
"""
        metadata, body = split_frontmatter(raw)
        self.assertEqual(metadata["doc_id"], "card-a")
        self.assertIn("# **测试套餐**", body)

        with TemporaryDirectory() as directory:
            path = Path(directory) / "card.md"
            path.write_text(raw, encoding="utf-8")
            document = DocumentParserRegistry().parse(path)
        self.assertEqual(document.doc_id, "card-a")
        self.assertEqual(document.title, "测试套餐")
        self.assertEqual(document.tags, ("联通", "四川"))

    def test_structure_first_then_size_fallback(self):
        document = RagDocument(
            doc_id="doc-a",
            title="套餐 A",
            content="# 套餐 A\n\n## 资费\n\n" + "每月包含通用流量。" * 80,
            source="test.md",
            category="product",
        )
        chunks = StructureAwareChunker(ChunkingConfig(180, 30)).split(document)
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk.content) <= 180 for chunk in chunks))
        self.assertTrue(all(chunk.section == "资费" for chunk in chunks))
        self.assertTrue(all(chunk.content.startswith("套餐 A\n资费") for chunk in chunks))

    def test_chunk_ids_do_not_depend_on_position(self):
        chunker = StructureAwareChunker(ChunkingConfig(256, 32))
        original = RagDocument(
            doc_id="stable",
            title="套餐",
            content="# 套餐\n## 资费\n30 元\n## 激活\n实名认证",
            source="test.md",
        )
        changed = RagDocument(
            doc_id="stable",
            title="套餐",
            content="# 套餐\n## 介绍\n长期套餐\n## 资费\n30 元\n## 激活\n实名认证",
            source="test.md",
        )
        first = {chunk.section: chunk.id for chunk in chunker.split(original)}
        second = {chunk.section: chunk.id for chunk in chunker.split(changed)}
        self.assertEqual(first["资费"], second["资费"])
        self.assertEqual(first["激活"], second["激活"])

    def test_markdown_headings_inside_code_fence_are_not_sections(self):
        document = RagDocument(
            doc_id="code-doc",
            title="接口说明",
            content="# 接口说明\n## 示例\n```markdown\n## 这不是章节\n```\n调用示例。",
            source="api.md",
        )
        chunks = StructureAwareChunker(ChunkingConfig(256, 32)).split(document)
        self.assertEqual({chunk.section for chunk in chunks}, {"示例"})
        self.assertIn("## 这不是章节", chunks[0].content)

    def test_html_parser_preserves_headings_and_drops_script(self):
        html = "<html><head><title>帮助中心</title><script>bad()</script></head>"
        html += "<body><h2>退款</h2><p>联系客服处理。</p></body></html>"
        with TemporaryDirectory() as directory:
            path = Path(directory) / "help.html"
            path.write_text(html, encoding="utf-8")
            document = DocumentParserRegistry().parse(path)
        self.assertEqual(document.title, "帮助中心")
        self.assertIn("## 退款", document.content)
        self.assertNotIn("bad()", document.content)

    def test_docx_and_xlsx_parsers(self):
        from docx import Document as WordDocument
        from openpyxl import Workbook

        with TemporaryDirectory() as directory:
            root = Path(directory)
            docx_path = root / "policy.docx"
            word = WordDocument()
            word.add_heading("售后政策", level=1)
            word.add_paragraph("激活后可联系客服处理。")
            word.save(docx_path)

            xlsx_path = root / "plans.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "套餐"
            sheet.append(["名称", "月租"])
            sheet.append(["畅享卡", 29])
            workbook.save(xlsx_path)

            registry = DocumentParserRegistry()
            docx_document = registry.parse(docx_path)
            xlsx_document = registry.parse(xlsx_path)

        self.assertIn("# 售后政策", docx_document.content)
        self.assertIn("激活后可联系客服处理。", docx_document.content)
        self.assertIn("## 套餐", xlsx_document.content)
        self.assertIn("畅享卡 | 29", xlsx_document.content)

    def test_ingestion_skips_unchanged_document(self):
        store = FakeStore()
        service = IngestionService(store=store)
        prepared = service.prepare_text(
            doc_id="manual-a",
            title="注销说明",
            content="可在运营商 APP 内申请注销。",
        )
        first = service.ingest(prepared)
        second = service.ingest(prepared)
        self.assertEqual(first.status, "indexed")
        self.assertEqual(second.status, "skipped")
        self.assertEqual(store.replace_calls, 1)
        self.assertEqual(service.delete("manual-a"), 1)


if __name__ == "__main__":
    unittest.main()
