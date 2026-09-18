"""多格式文档解析器。

PDF、DOCX、XLSX 使用可选依赖并延迟导入，因此仅处理 Markdown/TXT/HTML/CSV
时不会加载大型解析库。
"""

import csv
import io
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Dict, Iterable, Protocol, Tuple

import yaml

from app.rag.schemas import RagDocument, normalize_tags


class UnsupportedDocumentError(ValueError):
    pass


_FM_FENCE_RE = re.compile(
    r"^```(?:markdown|md)?\s*\n---\n(.*?)\n---\n```\s*\n?(.*)$", re.S | re.I
)
_FM_BARE_RE = re.compile(r"^---\n(.*?)\n---\n?(.*)$", re.S)
_H1_RE = re.compile(r"^#\s+(.+?)\s*$", re.M)


def split_frontmatter(text: str) -> Tuple[Dict[str, Any], str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    match = _FM_FENCE_RE.match(normalized) or _FM_BARE_RE.match(normalized)
    if not match:
        return {}, normalized
    metadata = yaml.safe_load(match.group(1)) or {}
    if not isinstance(metadata, dict):
        raise ValueError("frontmatter 必须是键值对象")
    return metadata, match.group(2)


def _clean_heading(value: str) -> str:
    return re.sub(r"[*_`#]", "", value).strip()


def _relative_source(path: Path) -> str:
    try:
        return path.resolve().relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return path.resolve().as_posix()


def _build_document(
    path: Path,
    content: str,
    metadata: Dict[str, Any],
    detected_title: str = "",
) -> RagDocument:
    doc_id = str(metadata.get("doc_id") or path.stem)
    title = str(metadata.get("title") or detected_title or path.stem)
    category = str(metadata.get("category") or "knowledge")
    source = str(metadata.get("source") or _relative_source(path)).replace("\\", "/")
    known = {"doc_id", "title", "category", "tags", "source"}
    extra = {key: value for key, value in metadata.items() if key not in known}
    return RagDocument(
        doc_id=doc_id,
        title=_clean_heading(title),
        content=content,
        source=source,
        category=category,
        tags=normalize_tags(metadata.get("tags")),
        metadata=extra,
    )


class DocumentParser(Protocol):
    suffixes: Tuple[str, ...]

    def parse(self, path: Path, metadata: Dict[str, Any] | None = None) -> RagDocument:
        ...


class MarkdownParser:
    suffixes = (".md", ".markdown")

    def parse(self, path: Path, metadata: Dict[str, Any] | None = None) -> RagDocument:
        frontmatter, body = split_frontmatter(path.read_text(encoding="utf-8-sig"))
        merged = {**frontmatter, **(metadata or {})}
        match = _H1_RE.search(body)
        title = _clean_heading(match.group(1)) if match else path.stem
        return _build_document(path, body, merged, title)


class TextParser:
    suffixes = (".txt",)

    def parse(self, path: Path, metadata: Dict[str, Any] | None = None) -> RagDocument:
        return _build_document(path, path.read_text(encoding="utf-8-sig"), metadata or {})


class CsvParser:
    suffixes = (".csv",)

    def parse(self, path: Path, metadata: Dict[str, Any] | None = None) -> RagDocument:
        text = path.read_text(encoding="utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
        content = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        return _build_document(path, content, metadata or {})


class _ReadableHtmlParser(HTMLParser):
    _ignored = {"script", "style", "noscript", "svg"}
    _blocks = {"p", "div", "section", "article", "li", "tr", "br"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title_parts: list[str] = []
        self._ignored_depth = 0
        self._in_title = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag in self._ignored:
            self._ignored_depth += 1
        if tag == "title":
            self._in_title = True
        if not self._ignored_depth:
            if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                self.parts.append("\n" + "#" * int(tag[1]) + " ")
            elif tag in self._blocks:
                self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag == "title":
            self._in_title = False
        if tag in self._ignored and self._ignored_depth:
            self._ignored_depth -= 1
        elif not self._ignored_depth and tag in self._blocks.union({f"h{i}" for i in range(1, 7)}):
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._ignored_depth:
            return
        clean = re.sub(r"\s+", " ", data).strip()
        if clean:
            self.parts.append(clean + " ")
            if self._in_title:
                self.title_parts.append(clean)

    @property
    def text(self) -> str:
        lines = (re.sub(r"[ \t]+", " ", line).strip() for line in "".join(self.parts).splitlines())
        return "\n".join(line for line in lines if line)

    @property
    def title(self) -> str:
        return " ".join(self.title_parts).strip()


class HtmlDocumentParser:
    suffixes = (".html", ".htm")

    def parse(self, path: Path, metadata: Dict[str, Any] | None = None) -> RagDocument:
        parser = _ReadableHtmlParser()
        parser.feed(path.read_text(encoding="utf-8-sig"))
        return _build_document(path, parser.text, metadata or {}, parser.title)


class PdfParser:
    suffixes = (".pdf",)

    def parse(self, path: Path, metadata: Dict[str, Any] | None = None) -> RagDocument:
        try:
            import pdfplumber
        except ImportError as exc:
            raise RuntimeError("解析 PDF 需要安装 pdfplumber") from exc
        pages = []
        with pdfplumber.open(path) as pdf:
            for index, page in enumerate(pdf.pages, 1):
                text = (page.extract_text() or "").strip()
                if text:
                    pages.append(f"## 第 {index} 页\n\n{text}")
        return _build_document(path, "\n\n".join(pages), metadata or {})


class DocxParser:
    suffixes = (".docx",)

    def parse(self, path: Path, metadata: Dict[str, Any] | None = None) -> RagDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("解析 DOCX 需要安装 python-docx") from exc
        document = Document(path)
        lines: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style else "").lower()
            match = re.search(r"heading\s*(\d+)", style)
            lines.append(f"{'#' * min(int(match.group(1)), 6)} {text}" if match else text)
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return _build_document(path, "\n\n".join(lines), metadata or {})


class XlsxParser:
    suffixes = (".xlsx",)

    def parse(self, path: Path, metadata: Dict[str, Any] | None = None) -> RagDocument:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("解析 XLSX 需要安装 openpyxl") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        parts: list[str] = []
        try:
            for sheet in workbook.worksheets:
                parts.append(f"## {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() if value is not None else "" for value in row]
                    if any(values):
                        parts.append(" | ".join(values))
        finally:
            workbook.close()
        return _build_document(path, "\n".join(parts), metadata or {})


class DocumentParserRegistry:
    def __init__(self, parsers: Iterable[DocumentParser] | None = None) -> None:
        self._parsers: Dict[str, DocumentParser] = {}
        defaults = (
            MarkdownParser(), TextParser(), CsvParser(), HtmlDocumentParser(),
            PdfParser(), DocxParser(), XlsxParser(),
        )
        for parser in parsers or defaults:
            for suffix in parser.suffixes:
                self._parsers[suffix.lower()] = parser

    @property
    def supported_suffixes(self) -> Tuple[str, ...]:
        return tuple(sorted(self._parsers))

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in self._parsers

    def parse(self, path: Path, metadata: Dict[str, Any] | None = None) -> RagDocument:
        parser = self._parsers.get(path.suffix.lower())
        if not parser:
            supported = ", ".join(self.supported_suffixes)
            raise UnsupportedDocumentError(f"不支持 {path.suffix or '无扩展名'} 文件；支持：{supported}")
        return parser.parse(path, metadata)
